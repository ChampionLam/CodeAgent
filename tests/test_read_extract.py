"""文档提取层：read_extract 分派 + read_file 接线 + OCR 兜底。

被测契约（见 python/read_extract.py 模块头）：
  * 标准库路（.ipynb/.docx/.xlsx）零可选依赖可用；
  * anydoc 路的格式（.pdf 等）装了 anydoc 就转 Markdown，没装给教学式错误；
  * 纯扫描件（无文本层）抛 NeedsOcrExtraction —— 绝不返回乱码；
  * read_file 命中文档走提取层，并把元信息（extractor/chars）带给模型；
  * OCR 兜底可注入（hosted / 本地视觉线 / 无），三条路都要如实回报。

零网络：anydoc 与 pymupdf 都在本地跑；OCR 的 chat 是假实现。
"""
from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import unicodedata
import unittest
import zipfile

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(REPO, "python"))

import attachments  # noqa: E402
import ocr_fallback  # noqa: E402
import read_extract  # noqa: E402
import tools  # noqa: E402

CJK_FONT = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
PHRASE = "通天河副本攻略"
SCAN_PHRASE = "扫描件正文（图片里的字）"


def _nfkc(text: str) -> str:
    """PDF 字体抽字可能落到 CJK 兼容区（如 U+F9xx），比较前统一 NFKC。"""
    return unicodedata.normalize("NFKC", text or "")


def _have(name: str) -> bool:
    try:
        __import__(name)
        return True
    except Exception:
        return False


class DocCase(unittest.TestCase):
    """所有用例共用的临时目录 + 各种 fixture 生成器。"""

    def setUp(self) -> None:
        self.dir = tempfile.mkdtemp(prefix="desk-read-extract-")
        self.ws = os.path.join(self.dir, "ws")
        os.makedirs(self.ws)
        tools.ReadFileTool.ocr_fallback = None

    def tearDown(self) -> None:
        tools.ReadFileTool.ocr_fallback = None
        shutil.rmtree(self.dir, ignore_errors=True)

    # ── fixtures ──────────────────────────────────────────────────────────
    def path(self, name: str) -> str:
        return os.path.join(self.dir, name)

    def write_text_file(self, name: str, text: str) -> str:
        p = self.path(name)
        with open(p, "w", encoding="utf-8") as fh:
            fh.write(text)
        return p

    def write_ipynb(self, name: str = "nb.ipynb") -> str:
        nb = {
            "cells": [
                {"cell_type": "markdown", "metadata": {},
                 "source": ["# 标题\n", "通天河副本攻略要点\n"]},
                {"cell_type": "code", "execution_count": 1, "metadata": {},
                 "source": ["print('第二阶段：躲水柱')\n"],
                 "outputs": [{"output_type": "stream", "name": "stdout",
                              "text": ["第二阶段：躲水柱\n"]}]},
            ],
            "metadata": {"language_info": {"name": "python"}},
            "nbformat": 4, "nbformat_minor": 5,
        }
        p = self.path(name)
        with open(p, "w", encoding="utf-8") as fh:
            json.dump(nb, fh, ensure_ascii=False)
        return p

    def write_docx(self, name: str = "doc.docx") -> str:
        """手写最小合法 docx（word/document.xml），不依赖 python-docx。"""
        xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:document xmlns:w="%s"><w:body>'
            "<w:p><w:r><w:t>通天河副本攻略</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>第一步</w:t><w:tab/><w:t>清小怪</w:t></w:r></w:p>"
            "</w:body></w:document>" % read_extract.NS_W
        )
        p = self.path(name)
        with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("word/document.xml", xml)
        return p

    def write_xlsx(self, name: str = "book.xlsx") -> str:
        """用 openpyxl 造一份真 xlsx（读写两端相互独立，便于验证）。"""
        from openpyxl import Workbook

        p = self.path(name)
        wb = Workbook()
        ws = wb.active
        ws.title = "攻略"
        ws["A1"] = "通天河"
        ws["B1"] = "BOSS 血量"
        ws["A2"] = 42
        wb.save(p)
        return p

    def write_text_pdf(self, name: str = "text.pdf") -> str:
        """有真文本层的中文 PDF（嵌入 CJK 字体，需系统有 Noto CJK）。"""
        import fitz

        p = self.path(name)
        doc = fitz.open()
        page = doc.new_page()
        page.insert_font(fontname="cjk", fontfile=CJK_FONT)
        page.insert_text((72, 120), PHRASE + "：先清小怪，再打BOSS",
                         fontname="cjk", fontsize=18)
        page.insert_text((72, 160), "第二阶段注意躲避水柱，站右边安全",
                         fontname="cjk", fontsize=14)
        doc.save(p)
        doc.close()
        return p

    def write_scan_pdf(self, name: str = "scan.pdf") -> str:
        """纯图片 PDF：整页只有一张图，没有文本层。"""
        import fitz
        from PIL import Image, ImageDraw

        img_path = self.path("scan-src.png")
        im = Image.new("RGB", (900, 300), "white")
        ImageDraw.Draw(im).rectangle([40, 40, 200, 200], fill="blue")
        im.save(img_path)
        p = self.path(name)
        doc = fitz.open()
        page = doc.new_page(width=900, height=300)
        page.insert_image(fitz.Rect(0, 0, 900, 300), filename=img_path)
        doc.save(p)
        doc.close()
        return p


# ── 分派与上限 ───────────────────────────────────────────────────────────

class DispatchTest(DocCase):
    def test_known_formats_are_extractable(self) -> None:
        for name in ("通天河攻略.pdf", "a.docx", "b.xlsx", "c.ipynb",
                     "d.rtf", "e.epub", "f.odt", "g.PDF"):
            self.assertTrue(read_extract.is_extractable_document(name), name)

    def test_plain_and_unknown_files_are_not_extractable(self) -> None:
        for name in ("a.txt", "a.py", "a.png", "a.json", "noext"):
            self.assertFalse(read_extract.is_extractable_document(name), name)

    def test_unsupported_extension_raises(self) -> None:
        """这条用例的本意是「读不了就报错、绝不返回乱码」。

        2026-09-26 起 .txt/.log 这类纯文本已能被读取（用户明确要求支持），
        所以改用真·二进制样本守同一件事。
        """
        p = os.path.join(self.ws, "blob.bin")
        with open(p, "wb") as fh:
            fh.write(bytes(range(256)) * 4)
        with self.assertRaises(read_extract.ExtractionError):
            read_extract.extract_document_text(p)

    def test_oversize_document_raises(self) -> None:
        p = self.write_docx()
        old = read_extract.MAX_DOCUMENT_BYTES
        read_extract.MAX_DOCUMENT_BYTES = 16
        try:
            with self.assertRaises(read_extract.ExtractionError) as ctx:
                read_extract.extract_document_text(p)
        finally:
            read_extract.MAX_DOCUMENT_BYTES = old
        self.assertIn("过大", str(ctx.exception))

    def test_missing_file_raises_extraction_error(self) -> None:
        with self.assertRaises(read_extract.ExtractionError):
            read_extract.extract_document_text(self.path("nope.docx"))


# ── 标准库三种格式 ───────────────────────────────────────────────────────

class StdlibExtractTest(DocCase):
    def test_ipynb_extracts_markdown_code_and_output(self) -> None:
        text = read_extract.extract_document_text(self.write_ipynb())
        self.assertIn("通天河副本攻略要点", text)
        self.assertIn("print('第二阶段：躲水柱')", text)
        self.assertIn("第二阶段：躲水柱", text)

    def test_docx_extracts_paragraph_text(self) -> None:
        text = read_extract.extract_document_text(self.write_docx())
        self.assertIn(PHRASE, text)
        self.assertIn("清小怪", text)

    @unittest.skipUnless(_have("openpyxl"), "openpyxl 未安装")
    def test_xlsx_extracts_text_and_numbers(self) -> None:
        text = read_extract.extract_document_text(self.write_xlsx())
        self.assertIn("攻略", text)          # 工作表名
        self.assertIn("通天河", text)
        self.assertIn("BOSS 血量", text)
        self.assertIn("42", text)

    def test_bad_zip_docx_reports_docx_error(self) -> None:
        p = self.write_text_file("fake.docx", "not a zip at all")
        with self.assertRaises(read_extract.ExtractionError) as ctx:
            read_extract.extract_document_text(p)
        self.assertIn("DOCX", str(ctx.exception))


# ── PDF：文本层 vs 扫描件 ────────────────────────────────────────────────

@unittest.skipUnless(_have("fitz"), "pymupdf 未安装")
class PdfExtractTest(DocCase):
    @unittest.skipUnless(os.path.exists(CJK_FONT), "系统缺 CJK 字体")
    def test_text_layer_pdf_extracts_chinese(self) -> None:
        text = _nfkc(read_extract.extract_document_text(self.write_text_pdf()))
        self.assertIn(PHRASE, text)
        self.assertIn("躲避水柱", text)

    def test_scanned_pdf_needs_ocr_instead_of_garbage(self) -> None:
        p = self.write_scan_pdf()
        with self.assertRaises(read_extract.NeedsOcrExtraction) as ctx:
            read_extract.extract_document_text(p)
        self.assertEqual(ctx.exception.pages, [1])

    def test_needs_ocr_message_is_actionable_and_has_no_phantom_tool(self) -> None:
        msg = read_extract.needs_ocr_message("/tmp/x.pdf", [1, 2])
        self.assertIn("OCR", msg)
        self.assertIn("1, 2", msg)
        # 曾经指向不存在的 read_image 工具 → 模型会去调、连失败撞失败上限。
        self.assertNotIn("read_image", msg)
        self.assertNotIn("vision_analyze", msg)


# ── read_file 接线 ──────────────────────────────────────────────────────

class ReadFileWiringTest(DocCase):
    def _run(self, path: str):
        return tools.ReadFileTool().run({"path": path}, workspace_root=self.ws)

    def test_plain_text_file_path_unchanged(self) -> None:
        p = self.write_text_file("a.py", "print('hi')\n")
        res = self._run(p)
        self.assertTrue(res.ok)
        self.assertIn("print('hi')", res.content)
        self.assertNotIn("extractor", res.data)

    def test_not_found_still_not_found(self) -> None:
        res = self._run(self.path("missing.pdf"))
        self.assertFalse(res.ok)
        self.assertEqual(res.error_code, "NOT_FOUND")

    @unittest.skipUnless(_have("fitz") and os.path.exists(CJK_FONT), "缺 pymupdf/字体")
    def test_pdf_is_extracted_not_decoded(self) -> None:
        res = self._run(self.write_text_pdf())
        self.assertTrue(res.ok, res.error_message)
        self.assertIn(PHRASE, _nfkc(res.content))
        self.assertEqual(res.data["extractor"], "document")
        self.assertEqual(res.data["chars"], len(res.content))
        self.assertNotIn("\ufffd", res.content)

    def test_extract_failure_is_a_truthful_tool_error(self) -> None:
        p = self.write_text_file("broken.xlsx", "definitely not a workbook")
        res = self._run(p)
        self.assertFalse(res.ok)
        self.assertEqual(res.error_code, "EXTRACTION_FAILED")
        self.assertIn("read_file:", res.error_message)

    @unittest.skipUnless(_have("fitz"), "pymupdf 未安装")
    def test_scanned_pdf_without_fallback_returns_ocr_prompt(self) -> None:
        res = self._run(self.write_scan_pdf())
        self.assertTrue(res.ok)
        self.assertEqual(res.data["ocr"], "none")
        self.assertEqual(res.data["extractor"], "needs-ocr")
        self.assertIn("OCR", res.content)

    @unittest.skipUnless(_have("fitz"), "pymupdf 未安装")
    def test_injected_fallback_text_is_returned(self) -> None:
        calls = []

        def fake(path, exc):
            calls.append((path, list(exc.pages)))
            return "## 第 1 页\n" + SCAN_PHRASE, {"route": "vision", "pages": [1]}

        tools.ReadFileTool.ocr_fallback = fake
        res = self._run(self.write_scan_pdf())
        self.assertTrue(res.ok, res.error_message)
        self.assertEqual(len(calls), 1)
        self.assertEqual(res.data["extractor"], "ocr-vision")
        self.assertEqual(res.data["ocr"], "vision")
        self.assertIn(SCAN_PHRASE, res.content)

    @unittest.skipUnless(_have("fitz"), "pymupdf 未安装")
    def test_fallback_route_none_falls_back_to_prompt(self) -> None:
        tools.ReadFileTool.ocr_fallback = lambda path, exc: ("", {"route": "none"})
        res = self._run(self.write_scan_pdf())
        self.assertTrue(res.ok)
        self.assertEqual(res.data["ocr"], "none")
        self.assertIn("OCR", res.content)

    @unittest.skipUnless(_have("fitz"), "pymupdf 未安装")
    def test_fallback_crash_is_reported_not_swallowed(self) -> None:
        def boom(path, exc):
            raise RuntimeError("vision line down")

        tools.ReadFileTool.ocr_fallback = boom
        res = self._run(self.write_scan_pdf())
        self.assertFalse(res.ok)
        self.assertEqual(res.error_code, "OCR_FAILED")
        self.assertIn("vision line down", res.error_message)


# ── OCR 兜底本身 ────────────────────────────────────────────────────────

@unittest.skipUnless(_have("fitz"), "pymupdf 未安装")
class OcrFallbackTest(DocCase):
    def _needs_ocr(self, path: str):
        try:
            read_extract.extract_document_text(path)
        except read_extract.NeedsOcrExtraction as exc:
            return exc
        self.fail("该文件应当被判为纯扫描件")

    def test_local_route_calls_chat_with_image_data_url(self) -> None:
        p = self.write_scan_pdf()
        exc = self._needs_ocr(p)
        seen = []

        def chat(messages):
            seen.append(messages)
            return "第 1 页的文字：" + SCAN_PHRASE

        text, meta = ocr_fallback.extract_with_fallback(
            p, exc=exc, hosted_cfg=None, chat=chat)
        self.assertEqual(meta["route"], "vision")
        self.assertIn(SCAN_PHRASE, text)
        self.assertIn("## 第 1 页", text)
        self.assertEqual(len(seen), 1)
        block = seen[0][0]["content"]
        kinds = [part["type"] for part in block]
        self.assertIn("image_url", kinds)
        self.assertTrue(block[1]["image_url"]["url"].startswith("data:image/png;base64,"))

    def test_no_route_available_returns_none_route(self) -> None:
        p = self.write_scan_pdf()
        exc = self._needs_ocr(p)
        text, meta = ocr_fallback.extract_with_fallback(
            p, exc=exc, hosted_cfg=None, chat=None)
        self.assertEqual(text, "")
        self.assertEqual(meta["route"], "none")
        self.assertEqual(meta["pages"], [1])

    def test_local_chat_failure_is_reported_per_page(self) -> None:
        p = self.write_scan_pdf()
        exc = self._needs_ocr(p)

        def chat(messages):
            raise RuntimeError("model 500")

        text, meta = ocr_fallback.extract_with_fallback(
            p, exc=exc, hosted_cfg=None, chat=chat)
        # 逐页标注失败，不静默丢页；仍算走了本地路（有产物）。
        self.assertEqual(meta["route"], "vision")
        self.assertIn("OCR 失败", text)

    def test_hosted_config_requires_key(self) -> None:
        self.assertFalse(ocr_fallback.HostedOcrConfig(api_key=None).enabled)
        self.assertFalse(ocr_fallback.HostedOcrConfig(api_key="  ").enabled)
        self.assertTrue(ocr_fallback.HostedOcrConfig(api_key="k").enabled)


# ── 附件：文档识别与上下文 ──────────────────────────────────────────────

class AttachmentDocumentTest(DocCase):
    def setUp(self) -> None:
        super().setUp()
        self.store = attachments.AttachmentStore(os.path.join(self.dir, "attach"))

    def test_pdf_bytes_are_sniffed_as_document(self) -> None:
        p = self.write_text_pdf() if (_have("fitz") and os.path.exists(CJK_FONT)) else None
        if p is None:
            self.skipTest("缺 pymupdf/字体")
        with open(p, "rb") as fh:
            data = fh.read()
        self.assertEqual(attachments.sniff_mime(data), "application/pdf")
        ref = self.store.save_bytes(data, mime="application/pdf",
                                    original_name="通天河攻略.pdf")
        self.assertTrue(ref.is_document)
        self.assertFalse(ref.is_image)

    def test_document_text_is_extracted_and_cached(self) -> None:
        if not (_have("fitz") and os.path.exists(CJK_FONT)):
            self.skipTest("缺 pymupdf/字体")
        with open(self.write_text_pdf(), "rb") as fh:
            data = fh.read()
        ref = self.store.save_bytes(data, mime="application/pdf",
                                    original_name="通天河攻略.pdf")
        text = self.store.extract_document_text(ref)
        self.assertIn(PHRASE, _nfkc(text))
        self.assertTrue(os.path.exists(self.store.extracted_text_path(ref)))
        self.assertEqual(self.store.extract_document_text(ref), text)

    def test_document_context_text_caps_length(self) -> None:
        if not (_have("fitz") and os.path.exists(CJK_FONT)):
            self.skipTest("缺 pymupdf/字体")
        with open(self.write_text_pdf(), "rb") as fh:
            data = fh.read()
        ref = self.store.save_bytes(data, mime="application/pdf",
                                    original_name="通天河攻略.pdf")
        capped = attachments.document_context_text(self.store, [ref], max_chars=40)
        self.assertLessEqual(len(capped), 200)

    def test_png_is_still_an_image_not_a_document(self) -> None:
        from PIL import Image

        p = self.path("x.png")
        Image.new("RGB", (20, 20), "red").save(p)
        with open(p, "rb") as fh:
            ref = self.store.save_bytes(fh.read(), mime="image/png",
                                        original_name="x.png")
        self.assertTrue(ref.is_image)
        self.assertFalse(ref.is_document)


if __name__ == "__main__":
    unittest.main()

class OcrThinkingStripTests(unittest.TestCase):
    """本地 OCR 绝不能把模型的思考块写进正文（.5 真机抓到过）。"""

    def test_strips_closed_thinking_block(self):
        from ocr_fallback import _strip_thinking
        raw = "## \u7b2c 1 \u9875\n<think\u003e\u63a8\u7406\u8fc7\u7a0b...\u003c/think\u003e\n7 3 9"
        self.assertEqual(_strip_thinking(raw).splitlines()[-1], "7 3 9")
        self.assertNotIn("think", _strip_thinking(raw))
        self.assertNotIn("\u63a8\u7406\u8fc7\u7a0b", _strip_thinking(raw))

    def test_strips_thinking_tag_variant(self):
        from ocr_fallback import _strip_thinking
        self.assertEqual(_strip_thinking("<thinking>abc</thinking>\u6b63\u6587"),
                         "\u6b63\u6587")

    def test_strips_unclosed_thinking_tail(self):
        from ocr_fallback import _strip_thinking
        self.assertEqual(_strip_thinking("\u6b63\u6587<think\u003e\u88ab\u622a\u65ad\u7684\u63a8\u7406"),
                         "\u6b63\u6587")

    def test_leaves_plain_text_alone(self):
        from ocr_fallback import _strip_thinking
        self.assertEqual(_strip_thinking("  \u7eaf\u6b63\u6587  "), "\u7eaf\u6b63\u6587")


def _scanned_pdf(pages: int):
    """造一份 N 页、纯图片（无文本层）的 PDF，用来测 OCR 分页截断。"""
    import fitz

    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page(width=200, height=200)
        page.draw_rect(fitz.Rect(20, 20, 120, 120), color=(0, 0, 1), fill=(0, 0, 1))
    path = tempfile.mkstemp(suffix=".pdf")[1]
    doc.save(path)
    doc.close()
    return path


class PageArgParseTests(unittest.TestCase):
    """read_file 的 pages 参数解析（大扫描件续读的入口）。"""

    def test_range(self):
        self.assertEqual(read_extract.parse_pages_arg("9-16"), list(range(9, 17)))

    def test_reversed_range(self):
        self.assertEqual(read_extract.parse_pages_arg("16-9"), list(range(9, 17)))

    def test_single_and_list(self):
        self.assertEqual(read_extract.parse_pages_arg("3"), [3])
        self.assertEqual(read_extract.parse_pages_arg("1,5,9"), [1, 5, 9])

    def test_chinese_comma_and_spaces(self):
        self.assertEqual(read_extract.parse_pages_arg(" 1\uff0c5 , 9 "), [1, 5, 9])

    def test_junk_is_dropped_not_raised(self):
        self.assertEqual(read_extract.parse_pages_arg("abc"), [])
        self.assertEqual(read_extract.parse_pages_arg(""), [])
        self.assertEqual(read_extract.parse_pages_arg(None), [])
        self.assertEqual(read_extract.parse_pages_arg("1, abc ,3"), [1, 3])

    def test_page_count_filter(self):
        self.assertEqual(read_extract.parse_pages_arg("1-20", page_count=5), [1, 2, 3, 4, 5])


class OcrPageCapTests(unittest.TestCase):
    """页数超上限不再放弃：先给前 N 页 + 告诉模型剩下的页怎么续读。"""

    def test_render_pages_truncates_and_reports(self):
        path = _scanned_pdf(12)
        limits = ocr_fallback.OcrLimits(max_pages=8)
        rendered, skipped = ocr_fallback._render_pages(path, list(range(1, 13)), limits)
        self.assertEqual([p for p, _ in rendered], list(range(1, 9)))
        self.assertEqual(skipped, [9, 10, 11, 12])
        self.assertTrue(all(isinstance(png, bytes) and png for _, png in rendered))

    def test_render_pages_no_truncation(self):
        path = _scanned_pdf(3)
        rendered, skipped = ocr_fallback._render_pages(path, [1, 2, 3], ocr_fallback.OcrLimits())
        self.assertEqual([p for p, _ in rendered], [1, 2, 3])
        self.assertEqual(skipped, [])

    def test_ocr_note_when_truncated(self):
        path = _scanned_pdf(12)
        seen = []

        def chat(messages):
            seen.append(messages)
            return "第X页正文"

        stats = {}
        text = ocr_fallback.local_ocr(path, list(range(1, 13)), chat=chat,
                                      limits=ocr_fallback.OcrLimits(max_pages=8),
                                      stats=stats)
        self.assertEqual(len(seen), 8, "只该调 8 次模型")
        self.assertEqual(stats["pages_read"], list(range(1, 9)))
        self.assertEqual(stats["pages_skipped"], [9, 10, 11, 12])
        self.assertIn("pages", text)
        self.assertIn("9", text)
        self.assertIn("12", text)
        self.assertIn("## 第 1 页", text)

    def test_fallback_pages_override_reads_requested_pages(self):
        path = _scanned_pdf(12)
        exc = read_extract.NeedsOcrExtraction(path, list(range(1, 13)))
        seen = []

        def chat(messages):
            seen.append(messages)
            return "正文"

        text, meta = ocr_fallback.extract_with_fallback(
            path, exc=exc, hosted_cfg=None, chat=chat,
            limits=ocr_fallback.OcrLimits(max_pages=8),
            pages_override=[9, 10])
        self.assertEqual(meta["route"], "vision")
        self.assertEqual(meta["pages_read"], [9, 10])
        self.assertEqual(meta["pages_skipped"], [])
        self.assertEqual(len(seen), 2)
        self.assertIn("## 第 9 页", text)


class ReadFilePagesArgTests(unittest.TestCase):
    """read_file 的 pages 参数要真的透到兜底，且 schema 里对模型可见。"""

    def test_schema_exposes_pages(self):
        props = tools.ReadFileTool.parameters["properties"]
        self.assertIn("pages", props)
        self.assertEqual(props["pages"]["type"], "string")

    def test_pages_arg_reaches_fallback(self):
        path = _scanned_pdf(12)
        got = {}

        def fake_fallback(target, exc, pages=None):
            got["pages"] = pages
            return "正文", {"route": "vision"}

        old = tools.ReadFileTool.ocr_fallback
        tools.ReadFileTool.ocr_fallback = staticmethod(fake_fallback)
        try:
            res = tools.ReadFileTool().run({"path": path, "pages": "9-12"},
                                           workspace_root="/")
        finally:
            tools.ReadFileTool.ocr_fallback = old
        self.assertTrue(res.ok, res.error_message)
        self.assertEqual(got["pages"], [9, 10, 11, 12])

    def test_no_pages_arg_passes_empty(self):
        path = _scanned_pdf(2)
        got = {}

        def fake_fallback(target, exc, pages=None):
            got["pages"] = pages
            return "正文", {"route": "vision"}

        old = tools.ReadFileTool.ocr_fallback
        tools.ReadFileTool.ocr_fallback = staticmethod(fake_fallback)
        try:
            tools.ReadFileTool().run({"path": path}, workspace_root="/")
        finally:
            tools.ReadFileTool.ocr_fallback = old
        self.assertEqual(got["pages"], [])


class TextFileTest(DocCase):
    """纯文本类附件（用户 2026-09-26 拖 .log 报「不支持」后补的能力）。

    契约：UTF-8 / GBK / UTF-16 都读得出；二进制不许当文本；超限拒绝；
    未知后缀按内容嗅探；.csv 在没装 anydoc 时回落文本；.ipynb 不被嗅探截胡。
    """

    def _write(self, name: str, data: bytes) -> str:
        path = os.path.join(self.ws, name)
        with open(path, "wb") as fh:
            fh.write(data)
        return path

    def test_utf8_log_is_read(self) -> None:
        path = self._write("app.log", "[2026-09-26 10:00] 启动完成\n".encode("utf-8"))
        out = read_extract.extract_document_text(path)
        self.assertIn("启动完成", out)
        self.assertTrue(out.endswith("\n"))

    def test_gbk_log_is_read(self) -> None:
        """Windows 上的日志常是 GBK —— 解不出来就是一屏乱码。"""
        path = self._write("gbk.log", "错误：连接超时\n第二行\n".encode("gbk"))
        out = read_extract.extract_document_text(path)
        self.assertIn("连接超时", out)
        self.assertNotIn("\ufffd", out)

    def test_utf16_bom_log_is_read(self) -> None:
        path = self._write("u16.log", "hello utf16\n".encode("utf-16"))
        self.assertIn("hello utf16", read_extract.extract_document_text(path))

    def test_binary_renamed_to_log_is_rejected(self) -> None:
        """改了后缀的 exe / 二进制不许当文本喂给模型。"""
        path = self._write("fake.log", bytes(range(0, 256)) * 4)
        with self.assertRaises(read_extract.ExtractionError) as cm:
            read_extract.extract_document_text(path)
        self.assertIn("二进制", str(cm.exception))

    def test_unknown_extension_text_gate(self) -> None:
        """附件闸门认内容：.trace 也该收（read_file 闸门不管它是另一回事）。"""
        path = self._write("server.trace", b"line one\nline two\n")
        self.assertTrue(read_extract.is_text_attachment(path))
        self.assertFalse(read_extract.is_extractable_document(path))
        self.assertIn("line two", read_extract.extract_document_text(path))

    def test_unknown_extension_binary_is_not_text(self) -> None:
        path = self._write("payload.weird", bytes(range(1, 256)) * 4)
        self.assertFalse(read_extract.is_text_attachment(path))

    def test_oversize_text_is_rejected(self) -> None:
        path = self._write("big.log", b"x" * 4096)
        old = read_extract.MAX_TEXT_BYTES
        read_extract.MAX_TEXT_BYTES = 1024
        self.addCleanup(lambda: setattr(read_extract, "MAX_TEXT_BYTES", old))
        with self.assertRaises(read_extract.ExtractionError) as cm:
            read_extract.extract_document_text(path)
        self.assertIn("上限", str(cm.exception))

    def test_output_is_truncated(self) -> None:
        path = self._write("long.log", b"y" * 5000)
        old = read_extract._MAX_OUTPUT_CHARS
        read_extract._MAX_OUTPUT_CHARS = 100
        self.addCleanup(lambda: setattr(read_extract, "_MAX_OUTPUT_CHARS", old))
        out = read_extract.extract_document_text(path)
        self.assertIn("已截断", out)
        self.assertLess(len(out), 400)

    def test_bytes_entry_reads_gbk(self) -> None:
        out = read_extract.extract_document_bytes("中文日志".encode("gbk"), "x.log")
        self.assertIn("中文日志", out)

    def test_csv_reads_without_anydoc(self) -> None:
        path = self._write("data.csv", "a,b\n1,2\n".encode("utf-8"))
        self.assertIn("a,b", read_extract.extract_document_text(path))

    def test_notebook_is_not_hijacked_by_text_sniff(self) -> None:
        """回归：.ipynb 本身是 JSON 文本，必须仍走 notebook 分支（字节版尤其）。"""
        path = self.write_ipynb()
        by_path = read_extract.extract_document_text(path)
        with open(path, "rb") as fh:
            data = fh.read()
        by_bytes = read_extract.extract_document_bytes(data, os.path.basename(path))
        self.assertTrue(by_path.strip())
        self.assertEqual(by_path, by_bytes)

    def test_log_is_a_document_attachment(self) -> None:
        """用户那次报的正是这条：拖 .log 进来必须被当可读附件收下。"""
        path = self._write("app.log", b"[10:00] boot ok\n")
        store = attachments.AttachmentStore(os.path.join(self.dir, "objects"))
        ref = store.save_file(path)
        self.assertEqual(ref.original_name, "app.log")
        self.assertTrue(ref.is_document)
        self.assertIn("boot ok",
                      read_extract.extract_document_bytes(
                          open(path, "rb").read(), "app.log"))

    def test_read_file_gate_unchanged(self) -> None:
        """回归护栏：纯文本不许进 read_file 的提取闸门（否则输出形态改了）。"""
        for name in ("x.log", "x.txt", "x.py", "x.csv", "x.md"):
            ext = os.path.splitext(name)[1]
            self.assertIn(ext, read_extract.STDLIB_TEXT_EXTENSIONS)
            self.assertNotIn(ext, read_extract.EXTRACTABLE_EXTENSIONS)
